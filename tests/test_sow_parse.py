"""Uploaded statements of work: text extraction, the rule-based parser, the
optional model path, POST /api/sow/parse and source_document in the SOW."""
from __future__ import annotations

import io
import json
from datetime import date

import pytest

from app.intake import sow_parse
from app.intake.sow_parse import SowParseError, heuristic_scope, parse_text, parse_upload

TODAY = date(2026, 9, 27)

MARKDOWN_SOW = """# Statement of Work: Payroll competitor map

## Objective
A ranked map of the 20 closest competitors to our payroll product, with pricing
and positioning for each, so the sales team can answer comparison questions.

## Milestones
1. Scope and sources - $600
   - Agreed list of competitors and sources
   - Research questions signed off
2. Findings draft - $1,000
   - Every claim links to a source
3. Final report (900 USDC)
   - Executive summary fits on one page

## Timeline
Final report due by 15 December 2026.

## Budget
Total fixed fee: $2,500
"""

PLAIN_SOW = """STATEMENT OF WORK
Acme Ltd - CSV export

OBJECTIVE
Add a CSV export to the reports page so finance can pull monthly numbers without
engineering help.

Milestone 1: Technical plan
Amount: USD 400
Acceptance criteria: Plan approved before any code is written; Export columns agreed

Milestone 2: Working build
Payment: 1,100 USDC
- Runs in a staging environment
- Automated tests pass

TIMELINE
The work should be completed within 6 weeks of kickoff.
"""

TABLE_SOW = """Project: Newsletter relaunch

Background: We want to relaunch our weekly newsletter and grow it to 5,000 subscribers.

Deliverables
| # | Deliverable | Acceptance | Fee |
|---|---|---|---|
| 1 | Audit and plan | Baseline metrics recorded | $1.2k |
| 2 | First campaigns live | Each campaign has a success metric; Weekly report | $2,000 |
| Total | | | $3,200 |

Acceptance criteria
- Open rate reported weekly
"""


def _pdf(lines: list[str]) -> bytes:
    """A one-page PDF with real text objects, built by hand (no writer library)."""
    ops = ["BT", "/F1 11 Tf", "14 TL", "72 760 Td"]
    for line in lines:
        esc = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({esc}) Tj T*")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def _docx() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading("Statement of Work: Support triage automation", level=1)
    doc.add_heading("Objective", level=2)
    doc.add_paragraph("Automate first-line support triage so every ticket is tagged and routed "
                      "within five minutes of arriving.")
    doc.add_heading("Milestones", level=2)
    table = doc.add_table(rows=3, cols=3)
    for row, cells in zip(table.rows, [("Milestone", "Acceptance criteria", "Amount"),
                                       ("Process map", "Current steps and owners written down", "$300"),
                                       ("Automation live", "Runs end to end on real tickets; Runbook written",
                                        "$700")]):
        for cell, text in zip(row.cells, cells):
            cell.text = text
    doc.add_heading("Deadline", level=2)
    doc.add_paragraph("Complete by 2026-11-30.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── rule-based parser ─────────────────────────────────────────────────────
def test_markdown_sow_fills_every_field():
    s = heuristic_scope(MARKDOWN_SOW, today=TODAY)
    assert s["method"] == "heuristic"
    assert s["outcome"].startswith("A ranked map of the 20 closest competitors")
    assert s["brief"].startswith("A ranked map") and len(s["brief"]) <= 141
    assert s["category_key"] == "research"
    assert [m["title"] for m in s["milestones"]] == ["Scope and sources", "Findings draft", "Final report"]
    assert [m["amount_usdc"] for m in s["milestones"]] == [600, 1000, 900]
    assert s["milestones"][0]["acceptance"] == ["Agreed list of competitors and sources",
                                                "Research questions signed off"]
    assert s["deadline"] == {"date": "2026-12-15", "mode": "date",
                             "text": "Final report due by 15 December 2026."}
    assert s["budget_usdc"] == 2500
    assert s["warnings"] == []


def test_plain_text_milestone_markers_and_duration():
    s = heuristic_scope(PLAIN_SOW, today=TODAY)
    assert s["title"] == "Acme Ltd - CSV export"
    assert s["outcome"].startswith("Add a CSV export") and "Milestone" not in s["outcome"]
    assert s["milestones"] == [
        {"title": "Technical plan", "amount_usdc": 400,
         "acceptance": ["Plan approved before any code is written", "Export columns agreed"]},
        {"title": "Working build", "amount_usdc": 1100,
         "acceptance": ["Runs in a staging environment", "Automated tests pass"]},
    ]
    # No budget line: the budget is the milestone sum, and says so.
    assert s["budget_usdc"] == 1500
    assert any("sum of the milestone amounts" in n for n in s["notes"])
    # "within 6 weeks" is counted from today, with a warning.
    assert s["deadline"]["date"] == "2026-11-08" and s["deadline"]["mode"] == "date"
    assert any("counted from today" in w for w in s["warnings"])


def test_table_milestones_k_amounts_and_job_criteria():
    s = heuristic_scope(TABLE_SOW, today=TODAY)
    assert s["title"] == "Newsletter relaunch"
    assert s["category_key"] == "growth"
    assert [(m["title"], m["amount_usdc"]) for m in s["milestones"]] == [
        ("Audit and plan", 1200), ("First campaigns live", 2000)]
    assert s["milestones"][1]["acceptance"] == ["Each campaign has a success metric", "Weekly report"]
    assert s["budget_usdc"] == 3200
    assert s["acceptance"] == ["Open rate reported weekly"]


def test_flattened_lists_from_pdfs_keep_criteria_under_their_milestone():
    """PDF text loses indentation: numbered milestones and "-" criteria end
    up at the same level, as do "-" milestones that carry amounts."""
    numbered = "Milestones\n1. Plan - $300\n- Plan approved\n2. Build - $700\n- Tests pass\n- Deployed\n"
    priced = "Milestones\n- Plan ($300)\n- Plan approved\n- Build ($700)\n- Tests pass\n"
    for text in (numbered, priced):
        s = heuristic_scope(text, today=TODAY)
        assert [(m["title"], m["amount_usdc"]) for m in s["milestones"]] == [("Plan", 300), ("Build", 700)]
        assert s["milestones"][0]["acceptance"] == ["Plan approved"]
        assert s["milestones"][1]["acceptance"][0] == "Tests pass"


def test_missing_fields_are_null_with_warnings():
    s = heuristic_scope("We need someone to tidy up our spreadsheet of leads.", today=TODAY)
    assert s["outcome"] == "We need someone to tidy up our spreadsheet of leads."
    assert s["milestones"] == [] and s["budget_usdc"] is None
    assert s["deadline"] == {"date": None, "mode": None, "text": None}
    warnings = " ".join(s["warnings"])
    for missing in ("No milestones", "No budget", "No deadline"):
        assert missing in warnings


def test_nothing_is_invented_for_an_empty_outline():
    s = heuristic_scope("## Objective\n\n## Budget\nTBD\n", today=TODAY)
    assert s["outcome"] is None and s["budget_usdc"] is None and s["category_key"] is None
    assert any("No objective" in w for w in s["warnings"])
    assert any("kind of work" in w for w in s["warnings"])


def test_past_deadline_is_dropped_and_ranges_use_the_upper_figure():
    s = heuristic_scope("Objective: Write five blog posts about payroll.\n"
                        "Deadline: 2025-01-31\nBudget: $400 - $600\n", today=TODAY)
    assert s["deadline"]["date"] is None and any("already passed" in w for w in s["warnings"])
    assert s["budget_usdc"] == 600 and any("range" in w for w in s["warnings"])


def test_budget_mismatch_is_flagged_not_fixed():
    text = MARKDOWN_SOW.replace("Total fixed fee: $2,500", "Total fixed fee: $3,000")
    s = heuristic_scope(text, today=TODAY)
    assert s["budget_usdc"] == 3000
    assert [m["amount_usdc"] for m in s["milestones"]] == [600, 1000, 900]
    assert any("add up to 2500" in w for w in s["warnings"])


def test_amounts():
    found = sow_parse.amounts("Pay $1,250.50, then USD 300, 2k USDC, $3k and 40 dollars; not 2026.")
    assert [str(a) for a in found] == ["1250.50", "300", "2000", "3000", "40"]


# ── files ─────────────────────────────────────────────────────────────────
def test_docx_upload_reads_headings_and_tables():
    data = _docx()
    s = parse_upload(data, "triage.docx", today=TODAY, use_llm=False)
    assert s["source"]["filename"] == "triage.docx" and len(s["source"]["sha256"]) == 64
    assert s["outcome"].startswith("Automate first-line support triage")
    assert s["category_key"] == "ops"
    assert s["milestones"] == [
        {"title": "Process map", "acceptance": ["Current steps and owners written down"], "amount_usdc": 300},
        {"title": "Automation live", "acceptance": ["Runs end to end on real tickets", "Runbook written"],
         "amount_usdc": 700},
    ]
    assert s["budget_usdc"] == 1000
    assert s["deadline"]["date"] == "2026-11-30"


def test_pdf_upload_extracts_text_and_pages():
    data = _pdf(["Objective", "Build a dashboard of weekly sales metrics from our SQL database.",
                 "Budget", "Total: $1,800 USDC", "Deadline: 2026-12-01"])
    s = parse_upload(data, "dash.pdf", today=TODAY, use_llm=False)
    assert s["source"]["pages"] == 1
    assert s["outcome"] == "Build a dashboard of weekly sales metrics from our SQL database."
    assert s["category_key"] == "data"
    assert s["budget_usdc"] == 1800
    assert s["deadline"]["date"] == "2026-12-01"


@pytest.mark.parametrize("name, data, code", [
    ("sow.exe", b"MZ", "UNSUPPORTED_FILE_TYPE"),
    ("sow.doc", b"x", "UNSUPPORTED_FILE_TYPE"),
    ("sow.txt", b"", "EMPTY_DOCUMENT"),
    ("sow.txt", b"a" * (sow_parse.MAX_BYTES + 1), "FILE_TOO_LARGE"),
    ("sow.pdf", b"%PDF-1.4 not really", "UNREADABLE_DOCUMENT"),
    ("sow.docx", b"PK not a zip", "UNREADABLE_DOCUMENT"),
    ("sow.txt", b"   \n\n", "EMPTY_DOCUMENT"),
])
def test_rejected_files(name, data, code):
    with pytest.raises(SowParseError) as err:
        parse_upload(data, name, use_llm=False)
    assert err.value.code == code


def test_filename_is_reduced_to_a_basename():
    s = parse_upload(b"Objective: Write a launch post for our API.", "../../etc/sow.md", use_llm=False)
    assert s["source"]["filename"] == "sow.md"


# ── model path ────────────────────────────────────────────────────────────
def _llm_reply(monkeypatch, reply):
    from app import llm
    calls = []

    def fake_chat(system, user, **kw):
        calls.append({"system": system, "user": user, **kw})
        if isinstance(reply, Exception):
            raise reply
        return reply
    monkeypatch.setattr(llm, "LLM_URL", "http://llm.test")
    monkeypatch.setattr(llm, "chat", fake_chat)
    return calls


def test_llm_answer_is_validated_and_unknown_amounts_dropped(monkeypatch):
    reply = json.dumps({
        "outcome": "A ranked map of the 20 closest payroll competitors.",
        "brief": "Map the 20 closest payroll competitors.",
        "category_key": "research",
        "milestones": [{"title": "Scope", "acceptance": ["Sources agreed"], "amount_usdc": 600},
                       {"title": "Report", "acceptance": "One-page summary", "amount_usdc": 4321}],
        "deadline": {"date": "2026-12-15", "mode": "date"},
        "budget_usdc": 2500,
    })
    calls = _llm_reply(monkeypatch, "```json\n" + reply + "\n```")
    s = parse_text(MARKDOWN_SOW, today=TODAY)
    assert calls and calls[0]["timeout"] == sow_parse.LLM_TIMEOUT
    assert s["method"] == "llm"
    assert s["brief"] == "Map the 20 closest payroll competitors."
    assert s["milestones"][0]["amount_usdc"] == 600
    assert s["milestones"][1] == {"title": "Report", "acceptance": ["One-page summary"], "amount_usdc": None}
    assert any("not in the document" in w for w in s["warnings"])
    assert s["budget_usdc"] == 2500 and s["deadline"]["date"] == "2026-12-15"


def test_llm_bad_category_and_nulls_fall_back_to_rules(monkeypatch):
    _llm_reply(monkeypatch, json.dumps({"outcome": None, "brief": None, "category_key": "legal",
                                        "milestones": [], "deadline": None, "budget_usdc": None}))
    s = parse_text(MARKDOWN_SOW, today=TODAY)
    assert s["method"] == "llm"
    assert s["category_key"] == "research"          # from the rules, not "legal"
    assert s["budget_usdc"] == 2500 and len(s["milestones"]) == 3


@pytest.mark.parametrize("reply", ["not json at all", RuntimeError("LLM unreachable: timed out"),
                                   json.dumps({"milestones": "three"})])
def test_llm_failure_uses_the_rules(monkeypatch, reply):
    _llm_reply(monkeypatch, reply)
    s = parse_text(MARKDOWN_SOW, today=TODAY)
    assert s["method"] == "heuristic"
    assert "parsed with rules instead" in s["notes"][0]
    assert s["budget_usdc"] == 2500


def test_no_llm_call_without_llm_url(monkeypatch):
    from app import llm
    monkeypatch.setattr(llm, "LLM_URL", "")
    monkeypatch.setattr(llm, "chat", lambda *a, **k: pytest.fail("model called without LLM_URL"))
    assert parse_text(MARKDOWN_SOW, today=TODAY)["method"] == "heuristic"


# ── API ───────────────────────────────────────────────────────────────────
def _upload(client, data: bytes, name: str, **headers):
    return client.post("/api/sow/parse", data={"file": (io.BytesIO(data), name)},
                       content_type="multipart/form-data", headers=headers)


def test_api_parses_a_file(client):
    resp = _upload(client, MARKDOWN_SOW.encode(), "competitors.md")
    assert resp.status_code == 200, resp.get_json()
    body = resp.get_json()
    assert body["source"]["filename"] == "competitors.md"
    assert body["source"]["sha256"] == __import__("hashlib").sha256(MARKDOWN_SOW.encode()).hexdigest()
    assert body["budget_usdc"] == 2500 and len(body["milestones"]) == 3
    assert body["text"].startswith("# Statement of Work")


def test_api_parses_json_text(client):
    resp = client.post("/api/sow/parse", json={"text": PLAIN_SOW})
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["source"]["filename"] == "pasted-text.txt" and body["budget_usdc"] == 1500


def test_api_rejects_bad_input(client):
    resp = _upload(client, b"MZ", "tool.exe")
    assert resp.status_code == 415 and resp.get_json()["code"] == "UNSUPPORTED_FILE_TYPE"
    resp = _upload(client, b"a" * (sow_parse.MAX_BYTES + 10), "big.txt")
    assert resp.status_code == 413 and resp.get_json()["code"] == "FILE_TOO_LARGE"
    resp = client.post("/api/sow/parse", data={}, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert client.post("/api/sow/parse", json={"nope": 1}).status_code == 400


def test_api_auth_matches_the_other_json_apis(client, monkeypatch):
    monkeypatch.setenv("MCP_API_TOKEN", "s3cret")
    data = MARKDOWN_SOW.encode()
    assert _upload(client, data, "a.md").status_code == 401
    assert _upload(client, data, "a.md", **{"Sec-Fetch-Site": "cross-site"}).status_code == 401
    assert _upload(client, data, "a.md", Origin="https://evil.example").status_code == 401
    assert _upload(client, data, "a.md", **{"Sec-Fetch-Site": "same-origin"}).status_code == 200
    assert _upload(client, data, "a.md", Authorization="Bearer s3cret").status_code == 200


def test_upload_is_not_written_to_disk(client, monkeypatch, tmp_path):
    from werkzeug import formparser

    def spooled(*a, **k):
        raise AssertionError("upload went through Werkzeug's temporary-file stream")
    monkeypatch.setattr(formparser, "SpooledTemporaryFile", spooled)
    big = (MARKDOWN_SOW + "\n" + "Filler line for size.\n" * 40_000).encode()
    assert len(big) > 600_000     # over Werkzeug's in-memory threshold
    assert _upload(client, big, "big.md").status_code == 200


# ── pages ─────────────────────────────────────────────────────────────────
def test_new_page_has_the_upload_zone_and_no_ctrl_enter_hint(client):
    html = client.get("/new").get_data(as_text=True)
    assert 'data-sow-drop' in html and "Upload SOW" in html
    assert "js/sow-upload.js" in html and "css/sow.css" in html
    assert "Ctrl</kbd>" not in html and "to continue." not in html


def test_home_and_job_form_offer_upload(client, db, agent):
    from app.models import Agent
    from app.seller.stamp import dev_stamp
    dev_stamp(db.session.get(Agent, agent))   # the job form is only offered for hireable agents
    db.session.commit()
    home = client.get("/").get_data(as_text=True)
    assert 'id="home-sow"' in home and "Or upload a statement of work" in home
    form = client.get(f"/jobs/new?agent={agent}").get_data(as_text=True)
    assert 'data-sow-drop' in form and 'name="source_sha256"' in form
